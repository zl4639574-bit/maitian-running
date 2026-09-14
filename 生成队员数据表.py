# -*- coding: utf-8 -*-
"""生成《队员数据表.xlsx》：①队员总表（含各项目最好成绩）②成绩明细 ③有成绩但不在名册里的人 ④说明
数据来源：data/team-data.js（原始提取）+ data/overrides.js（队长在网页上改的/加的）
"""
import io, json, os, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, u"队员数据表.xlsx")
EVENT_ORDER = [u'5000米', u'3000米', u'10000米', u'1500米', u'半马', u'全马', u'4公里', u'12公里', u'16公里']


def load_js(fn, var):
    p = os.path.join(HERE, fn)
    if not os.path.exists(p):
        return {}
    t = io.open(p, encoding='utf-8').read()
    return json.loads(t.split('=', 1)[1].strip().rstrip(';'))


def fmt(sec):
    if not sec:
        return ''
    sec = float(sec)
    h = int(sec // 3600)
    m = int((sec % 3600) // 60)
    s = sec - h * 3600 - m * 60
    if h:
        return '%d:%02d:%02d' % (h, m, round(s))
    return '%d:%02d' % (m, round(s))


def main():
    base = load_js('data/team-data.js', 'TEAM_DATA')
    ov = load_js('data/overrides.js', 'TEAM_OVERRIDES')

    hidden = set(ov.get('hidden') or [])
    edits = ov.get('memberEdits') or {}
    news = ov.get('newMembers') or []
    hidden_rec = set(ov.get('hiddenRecords') or [])

    # ---- 名册（含修改和新增）----
    members, seen_uids = [], set()
    for m in (base.get('roster') or []):
        e = edits.get(m['name']) or {}
        rec = dict(m)
        rec.update({k: v for k, v in e.items() if v not in (None, '')})
        rec['_from'] = u'原始名册'
        rec['_hidden'] = m['name'] in hidden
        members.append(rec)
    for m in news:
        r = dict(m)
        r['_from'] = u'队长新增'
        r['_hidden'] = False
        members.append(r)

    # ---- 成绩：所有比赛 + 自由成绩 ----
    records = []
    comp_recs = ov.get('compRecords') or {}
    for ds in (base.get('datasets') or []):
        rows = list(ds.get('records') or []) + list(comp_recs.get(ds.get('id')) or [])
        for r in rows:
            key = ds.get('id') + '|' + r.get('name', '') + '|' + str(r.get('sec'))
            if key in hidden_rec:
                continue
            records.append(dict(r, _src=ds.get('label', ''), _date=ds.get('date', ''), _kind=u'比赛/测速'))
    for c in (ov.get('competitions') or []):
        for r in (c.get('records') or []):
            records.append(dict(r, _src=c.get('name', ''), _date=c.get('date', ''), _kind=u'比赛/测速'))
    for r in (ov.get('results') or []):
        records.append(dict(r, _src=r.get('meet') or u'队伍上传', _date=r.get('date', ''), _kind=u'自由成绩'))

    # ---- 每人每项目最好成绩 ----
    best = {}
    for r in records:
        n, ev, sec = r.get('name'), r.get('event') or u'未标注', r.get('sec')
        if not n or not sec:
            continue
        k = (n, ev)
        if k not in best or sec < best[k]['sec']:
            best[k] = {'sec': sec, 'src': r.get('_src', ''), 'date': r.get('_date', '')}
    events = [e for e in EVENT_ORDER if any(ev == e for (_, ev) in best)]
    events += sorted(set(ev for (_, ev) in best) - set(events))

    from openpyxl import Workbook
    wb = Workbook()

    ws = wb.active
    ws.title = u'队员总表'
    head = [u'序号', u'姓名', u'性别', u'学院', u'专业', u'年级', u'身份', u'公开显示', u'成绩条数', u'来源'] \
        + [e + u'最好成绩' for e in events] + [u'备注标签']
    ws.append(head)
    by_name = {}
    for r in records:
        by_name.setdefault(r.get('name'), []).append(r)

    def sortkey(m):
        lv = m.get('level') or []
        return (0 if u'正式' in lv else (1 if u'预备' in lv else 2), m.get('name') or '')

    n_public = 0
    for i, m in enumerate(sorted(members, key=sortkey), 1):
        lv = m.get('level') or []
        pub = (u'正式' in lv or u'预备' in lv) and not m.get('_hidden')
        if pub:
            n_public += 1
        rs = by_name.get(m['name']) or []
        row = [i, m.get('name', ''), m.get('sex', ''), m.get('college', ''), m.get('major', ''),
               m.get('grade', ''), '/'.join(lv) if lv else u'未分级', u'是' if pub else u'否',
               len(rs), m.get('_from', '')]
        for e in events:
            b = best.get((m['name'], e))
            row.append(fmt(b['sec']) if b else '')
        row.append('/'.join(m.get('tags') or []))
        ws.append(row)
    ws.freeze_panes = 'C2'

    ws2 = wb.create_sheet(u'成绩明细')
    ws2.append([u'姓名', u'性别', u'项目', u'成绩', u'秒', u'日期', u'来源', u'类型', u'名次/备注'])
    for r in sorted(records, key=lambda x: (x.get('name') or '', x.get('event') or '')):
        ws2.append([r.get('name', ''), r.get('sex', ''), r.get('event', ''), r.get('fmt') or fmt(r.get('sec')),
                    r.get('sec'), r.get('_date', ''), r.get('_src', ''), r.get('_kind', ''),
                    (r.get('rank') or r.get('note') or '')])
    ws2.freeze_panes = 'A2'

    names_in_roster = set(m.get('name') for m in members)
    orphan = sorted(set(r.get('name') for r in records if r.get('name') and r.get('name') not in names_in_roster))
    ws3 = wb.create_sheet(u'有成绩但不在名册里')
    ws3.append([u'姓名', u'成绩条数', u'涉及项目'])
    for n in orphan:
        rs = by_name.get(n) or []
        ws3.append([n, len(rs), '/'.join(sorted(set(r.get('event') or '' for r in rs)))])

    ws4 = wb.create_sheet(u'说明')
    for line in [
        [u'队员数据表'],
        [u'生成时间', datetime.datetime.now().strftime('%Y-%m-%d %H:%M')],
        [u'数据来源', u'data/team-data.js（从原始 Excel/Word 提取）+ data/overrides.js（队长在网页上改的/加的）'],
        [u'名册人数', len(members)],
        [u'公开显示（正式/预备）', n_public],
        [u'成绩记录条数', len(records)],
        [u'有最好成绩的项目', '/'.join(events) if events else u'（无）'],
        [u''],
        [u'读法'],
        [u'· 队员总表：一人一行；某项目的最好成绩 = 该项目所有成绩里最快的一次（自动取，改了成绩表它会自己变）'],
        [u'· 成绩明细：一条成绩一行，便于核对哪一场、哪一天跑的'],
        [u'· “有成绩但不在名册里”：这些名字有成绩但名册里没有这个人，多半是原始名单漏录或姓名写法不一致'],
        [u'· 想改数据：名册信息在 队长版→数据管理→队员名册 改；成绩在 队长版→数据管理→比赛成绩 增删改，改完点同步'],
    ]:
        ws4.append(line)

    for ws_ in (ws, ws2, ws3):
        for col in ws_.columns:
            w = max(len(str(c.value or '')) for c in col)
            ws_.column_dimensions[col[0].column_letter].width = min(max(w * 1.6 + 2, 8), 22)

    wb.save(OUT)
    import shutil
    shutil.copyfile(OUT, os.path.join(os.path.dirname(HERE.rstrip('\\/')), u'队员数据表.xlsx'))
    print(u'已生成: %s' % OUT)
    print(u'  队员总表 %d 人（公开显示 %d 人）| 成绩明细 %d 条 | 有成绩但不在名册 %d 人'
          % (len(members), n_public, len(records), len(orphan)))
    print(u'  项目: %s' % ('/'.join(events) or u'无'))
    if orphan:
        print(u'  不在名册里的名字（前 10）:', '、'.join(orphan[:10]))
    # CSV 方便直接看
    import csv
    with io.open(os.path.join(HERE, u'队员数据表.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        wr = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            wr.writerow(row)
    print(u'  另存 CSV: 队员数据表.csv')


if __name__ == '__main__':
    main()
