# -*- coding: utf-8 -*-
"""
麦田守望长跑队 · 数据中心 —— 数据提取脚本
把「E:\Desktop\麦田」里散落的 Excel / Word 资料，汇总成网页用的 data/team-data.js

用法（资料更新后重跑）：
  C:\\Users\\15749\\.venvs\\bib\\Scripts\\python.exe 提取数据.py
"""
import os, re, json, datetime

import openpyxl
import xlrd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                       # E:\Desktop\麦田
OUT  = os.path.join(HERE, "data", "team-data.js")

WARN = []

# ---------------------------------------------------------------- 工具函数

def norm(s):
    s = str(s).strip()
    for a, b in [("’", "'"), ("‘", "'"), ("”", '"'), ("“", '"'), ("：", ":"),
                 ("′", "'"), ("″", '"'), ("　", ""), ("，", ",")]:
        s = s.replace(a, b)
    return s.strip()


NONRESULT = re.compile(r"(?i)dns|dnf|缺|误|请假|未参加|无成绩|退赛")


def parse_sec(v):
    """各种写法 -> 秒。返回 (秒, 备注)"""
    if v is None:
        return None, ""
    if isinstance(v, datetime.time):                       # Excel 时间格：时分秒实为分:秒
        return round(v.hour * 60 + v.minute + v.second / 60.0, 1), ""
    if isinstance(v, datetime.datetime):
        return None, "日期格式"
    if isinstance(v, datetime.timedelta):
        return round(v.total_seconds(), 1), ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 0 < v < 1:                                      # Excel 时间序列号
            return round(v * 86400, 1), ""
        if v < 60:                                         # 13.57 = 13.57 分钟
            return round(v * 60, 1), ""
        return float(v), ""                                # 已按秒

    s = norm(v)
    if not s:
        return None, ""
    if NONRESULT.search(s):
        return None, s

    note = ""
    m = re.search(r"[（(]\s*([0-9]+\s*k[mM]?|3k|5k|10k)\s*[)）]", s, re.I)
    if m:
        note = m.group(1)
    s2 = re.sub(r"[（(].*?[)）]", "", s).strip()

    if "'" in s2 or '"' in s2:                             # 分'秒"
        parts = [p.strip() for p in re.split(r"['\"]", s2) if p.strip() != ""]
        try:
            nums = [float(p) for p in parts]
        except ValueError:
            return None, s
        if len(nums) == 1:
            return nums[0], note
        if len(nums) == 2:
            return round(nums[0] * 60 + nums[1], 1), note
        return None, s

    if ":" in s2:
        try:
            nums = [float(p) for p in s2.split(":") if p.strip() != ""]
        except ValueError:
            return None, s
        if len(nums) == 3:
            if nums[2] == 0 and nums[0] < 60 and nums[1] < 60:   # 17:02:00 实为 17分02秒
                return round(nums[0] * 60 + nums[1], 1), note
            return round(nums[0] * 3600 + nums[1] * 60 + nums[2], 1), note
        if len(nums) == 2:
            return round(nums[0] * 60 + nums[1], 1), note
        return None, s

    try:
        f = float(s2)
    except ValueError:
        return None, s
    if f <= 0:
        return None, s
    return (round(f * 60, 1), note) if f < 60 else (f, note)


def fmt(sec):
    if sec is None:
        return "-"
    sec = float(sec)
    h, rest = divmod(sec, 3600)
    m, s = divmod(rest, 60)
    return ("%d:%02d:%02d" % (h, m, s)) if h >= 1 else ("%d:%02d" % (m, s))


def grade_of(stuid):
    m = re.match(r"^(20\d{2})", str(stuid or ""))
    return m.group(1) if m else ""


CHINESE_NAME = re.compile(r"^[\u4e00-\u9fa5·]{2,5}$")
BAD_NAMES = {"姓名", "序号", "合计", "无", "队员", "学院", "成绩", "单位", "备注",
             "男子", "女子", "男生", "女生", "总时长", "圈数"}


def clean(name):
    if name is None:
        return ""
    n = norm(name).replace(" ", "").replace("\n", "")
    n = re.sub(r"[（(].*?[)）]", "", n)          # 去掉（女）（领跑）之类
    n = re.sub(r"\d+$", "", n) if not re.match(r"^[\u4e00-\u9fa5·]+$", n) else n
    return n.strip()


def is_name(n):
    return bool(n) and n not in BAD_NAMES and bool(CHINESE_NAME.match(n))


# ---------------------------------------------------------------- 名册

members = {}
lookup = {}        # 姓名 -> 性别/学院（只作参考信息，不算队员来源）

def add_member(name, **kw):
    n = clean(name)
    if not is_name(n):
        return None
    m = members.setdefault(n, {"name": n, "sex": "", "college": "", "major": "",
                               "grade": "", "level": set(), "tags": set()})
    for k in ("sex", "college", "major", "grade"):
        if kw.get(k) and not m[k]:
            m[k] = str(kw[k]).strip()
    if kw.get("level"):
        m["level"].add(kw["level"])
    if kw.get("tag"):
        m["tags"].add(kw["tag"])
    return m


def add_lookup(name, **kw):
    n = clean(name)
    if not is_name(n):
        return
    d = lookup.setdefault(n, {})
    for k in ("sex", "college", "grade"):
        if kw.get(k) and not d.get(k):
            d[k] = str(kw[k]).strip()


def read_roster_xlsx(rel, level, sheets=None):
    p = os.path.join(ROOT, rel)
    if not os.path.exists(p):
        WARN.append("缺文件: " + rel)
        return
    wb = openpyxl.load_workbook(p, data_only=True)
    for ws in wb.worksheets:
        if sheets and ws.title not in sheets:
            continue
        lvl = sheets.get(ws.title, level) if sheets else level
        for r in ws.iter_rows(min_row=3, values_only=True):
            if not r or len(r) < 2:
                continue
            name = r[1] if len(r) > 1 else None
            n = clean(name)
            if not is_name(n):
                continue
            add_member(n, college=r[2] if len(r) > 2 else "", major=r[3] if len(r) > 3 else "",
                       grade=grade_of(r[4]) if len(r) > 4 else "", level=lvl)


def read_roster():
    read_roster_xlsx(r"资料\队员信息\2023年在校正式队员信息表.xlsx", "正式")
    read_roster_xlsx(r"资料\队员信息\2023年预备队员信息表.xlsx", "预备")
    read_roster_xlsx(r"资料\队员信息\2023年预备队员信息表1.xlsx", "预备")
    # 训练考核表：两个 sheet 就是正式/预备名单
    read_roster_xlsx(r"训练统计\训练考核表.xlsx", "正式",
                     sheets={"正式队员": "正式", "预备队员": "预备"})
    # 队服 / 团服收集 = 订了队服的在队成员
    for rel in [r"队服信息收集.xlsx", r"团服尺码收集.xlsx",
                r"后稷跑协\后稷跑步协会团服（收集结果）.xlsx"]:
        p = os.path.join(ROOT, rel)
        if not os.path.exists(p):
            continue
        wb = openpyxl.load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(min_row=2, values_only=True):
                if r and r[0] is not None:
                    add_member(r[0], tag="队服")


def read_squads():
    """高百阵容 / 冬训营 / 社团骨干"""
    # 2025 高百资格赛两队名单（.xls）
    for rel, tag in [(r"高百\2025高百\5.17资格赛\2025年高百资格赛成绩统计1队.xls", "2025高百1队"),
                     (r"高百\2025高百\5.17资格赛\2025年高百资格赛成绩统计2队.xls", "2025高百2队")]:
        p = os.path.join(ROOT, rel)
        if not os.path.exists(p):
            continue
        try:
            wb = xlrd.open_workbook(p)
            for ws in wb.sheets():
                for i in range(ws.nrows):
                    raw = str(ws.cell_value(i, 0))
                    n = clean(raw)
                    if is_name(n):
                        add_member(n, tag=tag, sex="女" if "女" in raw else "")
                        if "领跑" in raw:
                            members[n]["tags"].add("领跑员")
        except Exception as e:
            WARN.append("读不了 %s: %s" % (rel, e))
    # 2024 高百甲组
    p = os.path.join(ROOT, r"高百\2024高百\5.19资格赛\甲组.xlsx")
    if os.path.exists(p):
        wb = openpyxl.load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            sex = "女" if "女" in ws.title else ("男" if "男" in ws.title else "")
            for r in ws.iter_rows(min_row=2, values_only=True):
                if r and r[0]:
                    add_member(r[0], tag="2024高百甲组", sex=sex)
    # 2025 届负责人及骨干
    leaders = [("李志宏", "负责人"), ("张津浩", "骨干"), ("轩辕馨玉", "骨干"),
               ("高宇航", "骨干"), ("李祥", "骨干"), ("董达运", "骨干")]
    for n, role in leaders:
        add_member(n, tag="25届" + role)


def read_lookup():
    """性别 / 学院参照表：越野赛与训练营名单（这些人不一定是队员，只借信息）"""
    p = os.path.join(ROOT, r"麦田活动\越野赛参赛名单（1）.xls")
    if os.path.exists(p):
        try:
            wb = xlrd.open_workbook(p)
            for ws in wb.sheets():
                if "教工" in ws.name:
                    continue
                for i in range(ws.nrows):
                    vals = [str(v).strip() for v in ws.row_values(i)]
                    if len(vals) < 5:
                        continue
                    sid, college, name, sex = vals[1], vals[2], vals[3], vals[4]
                    add_lookup(name, college=college, sex=sex, grade=grade_of(sid))
        except Exception as e:
            WARN.append("越野赛名单读取失败: %s" % e)


# ---------------------------------------------------------------- 成绩

datasets = []

def add_dataset(did, label, date, source, note="", event=""):
    d = {"id": did, "label": label, "date": date, "source": source,
         "note": note, "event": event, "records": []}
    datasets.append(d)
    return d


def push(ds, name, sex, sec, raw, note="", college="", event=""):
    n = clean(name)
    if sec is None or not is_name(n):
        return
    ds["records"].append({"name": n, "sex": sex, "sec": round(float(sec), 1),
                          "raw": norm(raw), "note": note, "college": college,
                          "event": event})
    add_member(n, sex=sex, college=college)


def sheet_rows(path, r0=1):
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)]
        head = " ".join(norm(c) for r in rows[:2] for c in r if isinstance(c, str))
        yield ws.title, head, rows[r0 - 1:]


def cells_of(r):
    return [c for c in r if c not in (None, "")]


def ds_10k():
    rel = r"训练统计\10.14——10.15测速.xlsx"
    p = os.path.join(ROOT, rel)
    if not os.path.exists(p):
        return
    ds = add_dataset("10k-2022", "10 公里测速", "2022.10.14—10.15", rel,
                     "男女分组路跑测速，DNS / DNF 未计成绩", "10000米")
    for title, head, rows in sheet_rows(p, 3):
        sex = "男" if title.startswith("男") else ("女" if title.startswith("女") else "")
        if not sex:
            continue
        for r in rows:
            cs = cells_of(r)
            if len(cs) < 2:
                continue
            n, val = clean(cs[-2]), cs[-1]
            sec, note = parse_sec(val)
            push(ds, n, sex, sec, val, note, event="10000米")


def ds_track(rel, did, date, label):
    p = os.path.join(ROOT, rel)
    if not os.path.exists(p):
        return
    ds = add_dataset(did, label, date, rel, "男生 5000 米 / 女生 3000 米，同一表内可能夹其他距离", "男女分组")
    for title, head, rows in sheet_rows(p, 1):
        sex = ""
        if "女" in title + head:
            sex = "女"
        elif "男" in title + head:
            sex = "男"
        if not sex:
            continue
        event = "3000米" if sex == "女" else "5000米"
        for r in rows:
            cs = cells_of(r)
            if not cs:
                continue
            # 表中间夹的小标题行，例如 “1500成绩” / “男生5000米”
            if len(cs) == 1 and isinstance(cs[0], str) and re.search(r"成绩|米|km", norm(cs[0])):
                m = re.search(r"(\d{3,5})", norm(cs[0]))
                if m:
                    event = m.group(1) + "米"
                    continue
            if len(cs) < 2:
                continue
            sec, note = parse_sec(cs[1])
            if sec is None:
                continue
            ev = event
            m = re.fullmatch(r"(?i)(\d+)\s*k", note or "")
            if m:                                    # 备注写了别的距离，按备注算
                ev = str(int(m.group(1)) * 1000) + "米"
            push(ds, cs[0], sex, sec, cs[1], note, event=ev)


def ds_winter():
    rel = r"冬训营\测速结果.xlsx"
    p = os.path.join(ROOT, rel)
    if not os.path.exists(p):
        return
    ds = add_dataset("winter-2022", "冬训营最终测速", "2022.12", rel,
                     "距离原表未标注；按性别推定：男生 5000 米 / 女生 3000 米", "")
    for title, head, rows in sheet_rows(p, 4):
        for r in rows:
            cs = cells_of(r)
            if len(cs) < 3:
                continue
            college, n, val = clean(cs[-3]), clean(cs[-2]), cs[-1]
            if not is_name(n):
                continue
            sex = lookup.get(n, {}).get("sex", "") or (members.get(n, {}) or {}).get("sex", "")
            sec, note = parse_sec(val)
            push(ds, n, sex, sec, val, note, college=college,
                 event=("3000米" if sex == "女" else ("5000米" if sex == "男" else "")))
            if n in members:
                members[n]["tags"].add("冬训营")


def ds_gaobai(rel, did, label, date, note, event):
    p = os.path.join(ROOT, rel)
    if not os.path.exists(p):
        return
    ds = add_dataset(did, label, date, rel, note, event)
    for title, head, rows in sheet_rows(p, 2):
        for r in rows:
            cs = cells_of(r)
            if len(cs) < 2:
                continue
            n, val = clean(cs[0]), cs[1]
            if not is_name(n):
                continue
            sec, note2 = parse_sec(val)
            push(ds, n, "", sec, val, note2, event=event)


# ---------------------------------------------------------------- 主流程

read_lookup()
read_roster()
read_squads()

ds_10k()
ds_track(r"训练统计\2024春季学期\3.31测速.xlsx", "track-20240331", "2024.3.31", "春季场地测速")
ds_track(r"训练统计\2024春季学期\4.28测速.xlsx", "track-20240428", "2024.4.28", "春季场地测速")
ds_winter()
ds_gaobai(r"高百\2023高百\2023.12.16总决赛\测速成绩.xlsx", "gb-20231216",
          "高百总决赛测速", "2023.12.16",
          "距离由原表配速反推（58'28\" ÷ 3'39\"/km ≈ 16km）", "16公里")
ds_gaobai(r"高百\2023高百\2023.9.29分站赛\赛前测速成绩.xlsx", "gb-20230929",
          "高百分站赛前测速", "2023.9.29",
          "距离由原表配速反推（≈4km）；原表成绩写作分钟小数", "4公里")

PB = [
    ("周金华", "男", ["5K 17:10", "10K 35:39", "半马 1:18", "全马 2:50"],
     "队伍创建者；2017 陕西省大学生运动会 10000 米甲组冠军"),
    ("龚才伟", "男", ["5K 16:41", "10K 34:59", "半马 1:15", "全马 2:47"],
     "已毕业；2022 贵州省全民运动会大众马拉松项目黔南代表"),
    ("马海云", "男", ["全马 2:46"], "国家二级运动员，多场马拉松官方配速员"),
    ("陈敏", "女", ["5K 16:30", "10K 34:03", "半马 1:13:23"],
     "2020 杨凌农科城马拉松大学生组半马第一"),
    ("阿巴小洛", "男", ["5K 16:45", "10K 34:50", "半马 1:16:42", "全马 2:53:22"],
     "2018 级，保研本校；2019 陕西省大学生田径运动会甲组第 7"),
    ("仇学铎", "男", ["5K 17:09", "10K 35:10", "半马 1:18:10", "全马 3:13:24"],
     "2019 校越野赛冠军；2021 杨凌马拉松大学生半程组第 4"),
]

# ---------------------------------------------------------------- 整理输出

roster = []
for m in members.values():
    lk = lookup.get(m["name"], {})
    for k in ("sex", "college"):
        if not m[k] and lk.get(k):
            m[k] = lk[k]
    if not m["grade"] and lk.get("grade"):
        m["grade"] = lk["grade"]
    m["level"] = sorted(m["level"]) or ["队员"]
    m["tags"] = sorted(m["tags"])
    m["inTeam"] = bool(m["level"] != ["队员"] or m["tags"])
    roster.append(m)
roster.sort(key=lambda x: (x["grade"] or "9999", x["name"]))

# 用全队信息补齐成绩记录里的性别 / 学院，并给没标距离的项目一个说法
sexmap = {m["name"]: m["sex"] for m in roster if m["sex"]}
colmap = {m["name"]: m["college"] for m in roster if m["college"]}
for ds in datasets:
    for r in ds["records"]:
        if not r["sex"] and r["name"] in sexmap:
            r["sex"] = sexmap[r["name"]]
        if not r["college"] and r["name"] in colmap:
            r["college"] = colmap[r["name"]]
        if not r["event"]:
            r["event"] = {"女": "3000米", "男": "5000米"}.get(r["sex"], "距离未标注")

for ds in datasets:
    # 同一数据集里可能夹着不同距离（如 5000 米表里混了 1500 米），按距离分组排名
    ds["records"].sort(key=lambda r: (r["event"] or "", r["sec"]))
    groups = {}
    for r in ds["records"]:
        groups.setdefault(r["event"] or ds["event"] or "成绩", []).append(r)
    for ev, recs in groups.items():
        recs.sort(key=lambda r: r["sec"])
        for i, r in enumerate(recs, 1):
            r["rank"] = i
            r["rankOf"] = ev
            r["fmt"] = fmt(r["sec"])
    ds["events"] = [{"event": ev, "count": len(recs),
                     "best": fmt(recs[0]["sec"]), "bestName": recs[0]["name"],
                     "worst": fmt(recs[-1]["sec"])}
                    for ev, recs in sorted(groups.items(), key=lambda kv: -len(kv[1]))]
    main = ds["events"][0] if ds["events"] else None
    ds["best"] = main["best"] if main else "-"
    ds["bestName"] = main["bestName"] if main else "-"
    ds["count"] = len(ds["records"])

memberset = {m["name"] for m in roster}
TEAM = {
    "name": "麦田守望长跑队",
    "alias": "西北农林科技大学 · 后稷跑步协会",
    "founded": "2017.09.23",
    "slogan": "麦田守望，一往无前",
    "intro": "麦田守望长跑队成立于 2017 年 9 月 23 日，聚集了西北农林科技大学所有爱跑步的"
             "校友和在校生，常年活跃在各种马拉松赛道上，代表学校参加省运会、大运会等中长跑"
             "项目以及高校百英里接力赛。",
    "media": [["微信公众号", "麦田守望长跑"], ["抖音", "mtsw_2017_CP"], ["B站", "麦田守望长跑队"]],
    "honors": [
        ["2017", "高校百英里接力赛西安站冠军 · 上海总决赛季军"],
        ["2018", "高校百英里接力赛西安站亚军，晋级总决赛"],
        ["2019", "高校百英里接力赛西安站第四，晋级总决赛"],
        ["2020", "高校百英里接力赛西安站季军（因疫情缺席总决赛）"],
        ["2021", "高校百英里接力赛西安站第四；杨凌国际马拉松大学生组男子包揽前八、女子包揽前三"],
        ["2022", "高校百英里资格赛西安站第一、全国第三；分站赛西安站季军；总决赛完赛"],
        ["2023", "杨凌农科城马拉松大学生组前二十占 14 人（前十占 8 人，含冠亚），女子包揽前四；高校百英里西安站第一、全国第三，分站赛第五，总决赛第 29 名"],
        ["2024", "高校百英里接力赛西安分站赛第五"],
    ],
    "activities": [
        "日常训练：渭河拉练、后山拉练、测速、间歇",
        "校外赛事：马拉松、高校百英里接力赛",
        "校内赛事：校运会、越野赛、新生越野赛领跑",
        "团建：毕业合影、聚餐、渭河拾荒、趣味对抗赛",
    ],
}

data = {
    "generated": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
    "team": TEAM,
    "roster": roster,
    "datasets": datasets,
    "pb": [{"name": n, "sex": s, "items": it, "note": nt} for n, s, it, nt in PB],
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    f.write("/* 本文件由「提取数据.py」自动生成；资料更新后重跑脚本即可 */\n")
    f.write("window.TEAM_DATA = ")
    json.dump(data, f, ensure_ascii=False, indent=1)
    f.write(";\n")

# ---------------------------------------------------------------- 自检
print("生成:", OUT)
team_n = sum(1 for m in roster if m["inTeam"])
print("名册 %d 人（其中有明确在队来源的 %d 人）| 有学院 %d | 有性别 %d"
      % (len(roster), team_n,
         sum(1 for m in roster if m["college"]), sum(1 for m in roster if m["sex"])))
print("成绩数据集 %d 组:" % len(datasets))
for ds in datasets:
    evs = " / ".join("%s %d条 最快%s(%s)" % (e["event"], e["count"], e["best"], e["bestName"])
                     for e in ds["events"])
    print("  %-14s %-14s 共%3d条  %s" % (ds["label"], ds["date"], ds["count"], evs))
if WARN:
    print("警告:")
    for w in WARN:
        print("  -", w)
