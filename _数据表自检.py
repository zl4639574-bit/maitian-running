# -*- coding: utf-8 -*-
"""验证三件事：
   ① 个人最好成绩只算正式队员（非正式同学只在单场榜里）
   ② 「导出队员数据表」能导出正确的 Excel（用 openpyxl 检查内容）
   ③ 「导入修正表」：改队员信息 / 改成绩 / 删成绩 / 补录成绩 → 预览差异 → 应用（全程只在本机，不写线上）
"""
import base64, importlib.util, io, json, os, re, sqlite3, subprocess, shutil, time, urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
OWNER, REPO = "zl4639574-bit", "maitian-running"
SITE = "https://%s.github.io/%s/" % (OWNER, REPO)
PAY = os.path.join(os.environ.get("LOCALAPPDATA", "."), "Temp", "mt_payload.txt")
TMP = os.path.join(os.environ.get("LOCALAPPDATA", "."), "Temp", "mt_dt")

spec = importlib.util.spec_from_file_location("ph", os.path.join(HERE, "手机检查.py"))
ph = importlib.util.module_from_spec(spec); spec.loader.exec_module(ph)


def payload():
    if os.path.exists(PAY):
        return io.open(PAY, encoding="utf-8").read().strip()
    db = os.path.join(os.environ["LOCALAPPDATA"], "hermes", "state.db").replace("\\", "/")
    cn = sqlite3.connect("file:%s?mode=ro" % db, uri=True)
    for mid, content in cn.execute("select id, content from messages where content like '%index.html#t=%' order by id desc limit 5"):
        for m in re.finditer(r"#t=([A-Za-z0-9_\-]+)", content or ""):
            p = m.group(1)
            try:
                o = json.loads(base64.urlsafe_b64decode(p + "=" * (-len(p) % 4)))
            except Exception:
                continue
            if o.get("token") and o.get("repo"):
                io.open(PAY, "w", encoding="utf-8").write(p)
                return p
    raise SystemExit("没载荷")


def tab_js(label, wait=1.4):
    return ("(async () => { const wait = ms => new Promise(r => setTimeout(r, ms));"
            " for (const n of document.querySelectorAll('.nav-item,[data-tab]'))"
            "  if ((n.textContent||'').trim() === %s) { n.click(); break; }"
            " await wait(%d); return 'ok'; })()" % (json.dumps(label), int(wait * 1000)))


def main():
    os.makedirs(TMP, exist_ok=True)
    pl = payload()
    port = ph.free_port(); proc, prof = ph.start_edge(port)
    try:
        time.sleep(2)
        with urllib.request.urlopen("http://127.0.0.1:%d/json/list" % port, timeout=8) as r:
            pg = next(t for t in json.loads(r.read()) if t["type"] == "page")
        c = ph.CDP(pg["webSocketDebuggerUrl"])
        for d in ("Page.enable", "Runtime.enable", "Network.enable", "DOM.enable"):
            c.send(d)
        c.send("Network.setCacheDisabled", cacheDisabled=True)
        c.send("Emulation.setDeviceMetricsOverride", width=390, height=844, deviceScaleFactor=2, mobile=True)

        def load(i=0):
            for _ in range(3):
                ph.goto(c, "%scaptain/?r=%d%d#t=%s" % (SITE, int(time.time()), i, pl), wait=30)
                if (c.js("location.href") or "").startswith("https://"):
                    break
            for _ in range(40):
                time.sleep(0.5)
                try:
                    if int(str(c.js("document.querySelectorAll('.nav-item').length") or 0)) >= 3:
                        return
                except Exception:
                    pass

        load()
        # 清掉本机残留
        c.js("(function(){localStorage.removeItem('mt_ov_local_v1');return 1;})()")
        load(1)

        print("═══ ① 个人最好成绩是否只算正式队员 ═══")
        c.js(tab_js(u'成绩榜'))
        r = c.js(r"""(async () => {
  const wait = ms => new Promise(r => setTimeout(r, ms));
  const rows = () => [].map.call(document.querySelectorAll('table.tbl tbody tr'), e => e.innerText.replace(/\s+/g,' ').trim());
  const pb = rows();
  const out = {个人最好成绩条数: (document.body.innerText.match(/共\s*(\d+)\s*条/) || ['','?'])[1],
               正式队员样例: pb.filter(x => x.indexOf('阿巴小洛') >= 0).slice(0, 3),
               非正式同学是否在最好成绩榜: pb.filter(x => x.indexOf('吕佳悦') >= 0 || x.indexOf('叶佳成') >= 0)};
  // 切到某场比赛（冬训营最终测速）
  const chips = [].map.call(document.querySelectorAll('.chip[data-comp]'), e => ({ id: e.dataset.comp, t: e.textContent.replace(/\s+/g,' ').trim() }));
  const tgt = chips.find(x => x.t.indexOf('冬训营') >= 0);
  out.比赛胶囊 = chips.map(x => x.t).slice(0, 8);
  if (tgt) { document.querySelector('.chip[data-comp="' + tgt.id + '"]').click(); await wait(1800); }
  const comp = rows();
  out.单场榜里非正式同学 = comp.filter(x => x.indexOf('吕佳悦') >= 0 || x.indexOf('叶佳成') >= 0).slice(0, 4);
  out.提示语 = (document.body.innerText.match(/个人最好成绩只统计[^\n]{0,40}/) || ['(没有提示语)'])[0];
  return JSON.stringify(out);
})()""")
        print("   ", r)

        print("\n═══ ② 导出队员数据表（捕获真实生成的 Excel 检查内容）═══")
        c.js(tab_js(u'数据管理'))
        c.js("(function(){var ch=document.querySelector('[data-msec=\"table\"]'); if(ch) ch.click(); return 1;})()")
        time.sleep(1)
        names = c.js(r"""(async () => {
  const wait = ms => new Promise(r => setTimeout(r, ms));
  window.__wb = null;
  const orig = XLSX.writeFile;
  XLSX.writeFile = function (wb) { window.__wb = wb; };
  const b = document.getElementById('btnDtExport');
  if (!b) return '没有导出按钮';
  b.click(); await wait(1500);
  XLSX.writeFile = orig;
  return window.__wb ? window.__wb.SheetNames.join(' / ') : '没生成工作簿';
})()""")
        print("    工作表:", names)
        b64 = c.js("(function(){ if(!window.__wb) return ''; return XLSX.write(window.__wb, {type:'base64', bookType:'xlsx'}); })()")
        xp = os.path.join(HERE, u"队员数据表_网页导出.xlsx")
        if b64:
            io.open(xp, "wb").write(base64.b64decode(b64))
            from openpyxl import load_workbook
            wb = load_workbook(xp)
            ws = wb[u'队员总表']; ws2 = wb[u'成绩明细']
            print("    导出文件:", os.path.basename(xp), "| 队员总表 %d 行 x %d 列 | 成绩明细 %d 行"
                  % (ws.max_row, ws.max_column, ws2.max_row))
            print("    总表表头:", [c.value for c in ws[1]][:11])
            print("    明细表头:", [c.value for c in ws2[1]])
        else:
            print("    ❌ 没拿到文件内容")

        print("\n═══ ③ 导入修正表（改信息 / 改成绩 / 删成绩 / 补录）═══")
        from openpyxl import load_workbook
        wb = load_workbook(xp)
        ws, ws2 = wb[u'队员总表'], wb[u'成绩明细']
        head = [c.value for c in ws[1]]
        ci = lambda name: head.index(name) + 1
        head2 = [c.value for c in ws2[1]]
        ci2 = lambda name: head2.index(name) + 1
        changed = []
        # (a) 改一个队员的学院
        for row in ws.iter_rows(min_row=2):
            if row[ci(u'姓名') - 1].value == u'郭家俊':
                row[ci(u'学院') - 1].value = u'动物科技学院（修正）'
                changed.append('郭家俊 学院→动物科技学院（修正）')
                break
        # (b) 改一条成绩：找 阿巴小洛 5000米 17:02 的那条
        target = None
        for row in ws2.iter_rows(min_row=2):
            if row[ci2(u'姓名') - 1].value == u'阿巴小洛' and row[ci2(u'项目') - 1].value == u'5000米':
                target = row; break
        if target:
            old_fmt = target[ci2(u'成绩') - 1].value
            target[ci2(u'成绩') - 1].value = '16:11'
            target[ci2(u'秒') - 1].value = 971
            changed.append(u'阿巴小洛 5000米 %s→16:11' % old_fmt)
        # (c) 删一条成绩：把某人某条成绩清空
        for row in ws2.iter_rows(min_row=2):
            if row[ci2(u'姓名') - 1].value == u'李娜':
                row[ci2(u'成绩') - 1].value = ''
                row[ci2(u'秒') - 1].value = None
                changed.append(u'李娜 删除一条成绩')
                break
        # (d) 补录一条成绩（编号留空）
        ws2.append(['', u'自检补录员', u'男', u'5000米', '15:59', 959, '2026-09-14', u'春季场地测速', u'比赛/测速', u'自检'])
        changed.append(u'补录 自检补录员 5000米 15:59')
        # (e) 新增一名队员
        ws.append([ws.max_row, u'自检新同学', u'女', u'林学院', u'林学2109', u'2024', u'正式', u'', '', u'', u'', u'', u'', u''])
        changed.append(u'新增队员 自检新同学')
        fx = os.path.join(TMP, u"修正示例.xlsx")
        wb.save(fx)
        print("    已构造修正文件:", os.path.basename(fx))
        for x in changed:
            print("      ·", x)

        doc = c.send("DOM.getDocument", depth=1)
        node = c.send("DOM.querySelector", nodeId=doc["root"]["nodeId"], selector="#dtFile")
        c.send("DOM.setFileInputFiles", files=[fx], nodeId=node["nodeId"])
        time.sleep(4)
        prev = c.js("(function(){var b=document.getElementById('dtBox');return b?b.innerText.replace(/\\s+/g,' ').slice(0,520):'(没有预览)';})()")
        print("    差异预览:", prev)

        print("    点「应用这些修正」…")
        r3 = c.js(r"""(async () => {
  const wait = ms => new Promise(r => setTimeout(r, ms));
  const b = document.getElementById('btnDtApply');
  if (!b) return '没有应用按钮';
  b.click(); await wait(1600);
  const l = JSON.parse(localStorage.getItem('mt_ov_local_v1') || '{}');
  return JSON.stringify({
    提示: (document.body.innerText.match(/已应用[^\n]{0,50}/) || ['(无)'])[0],
    改过的队员: Object.keys(l.memberEdits || {}),
    郭家俊新学院: (l.memberEdits || {})['郭家俊'] ? (l.memberEdits['郭家俊'].college || '') : '',
    隐藏的成绩条数: (l.hiddenRecords || []).length,
    新增的比赛成绩: Object.keys(l.compRecords || {}).map(k => (l.compRecords[k] || []).length),
    新增队员: (l.newMembers || []).map(m => m.name),
    自由成绩: (l.results || []).length});
})()""")
        print("   ", r3)

        print("\n    重新加载，看榜单/名册是否按修正后的数据变")
        load(2)
        r4 = c.js(r"""(async () => {
  const wait = ms => new Promise(r => setTimeout(r, ms));
  for (const n of document.querySelectorAll('.nav-item,[data-tab]'))
    if ((n.textContent||'').trim() === '成绩榜') { n.click(); break; }
  await wait(1800);
  const pb = [].map.call(document.querySelectorAll('table.tbl tbody tr'), e => e.innerText.replace(/\s+/g,' ').trim());
  const out = {阿巴小洛5000米: pb.filter(x => x.indexOf('阿巴小洛') >= 0 && x.indexOf('5000米') >= 0),
               自检补录员: pb.filter(x => x.indexOf('自检补录员') >= 0),
               还显示17:02吗: pb.some(x => x.indexOf('阿巴小洛') >= 0 && x.indexOf('17:02') >= 0)};
  for (const n of document.querySelectorAll('.nav-item,[data-tab]'))
    if ((n.textContent||'').trim() === '队员名册') { n.click(); break; }
  await wait(1600);
  const q = document.getElementById('rosterQ');
  if (q) { q.value = '自检新同学'; q.dispatchEvent(new Event('input', {bubbles: true})); await wait(1200); }
  out.新队员卡片 = (document.querySelector('.pcard') || {}).innerText ? document.querySelector('.pcard').innerText.replace(/\s+/g,' ').slice(0,90) : '(没找到)';
  return JSON.stringify(out);
})()""")
        print("   ", r4)

        sc = c.send("Page.captureScreenshot", format="png")
        io.open(os.path.join(HERE, u"数据表_验证.png"), "wb").write(base64.b64decode(sc["data"]))
        c.js("(function(){localStorage.removeItem('mt_ov_local_v1');return 1;})()")
        c.ws.close()
    finally:
        subprocess.run(["taskkill", "/F", "/PID", str(proc.pid), "/T"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        shutil.rmtree(prof, ignore_errors=True)
    print("\n（全程只写本机，测试数据已清空，线上未被改动）")


if __name__ == "__main__":
    main()
